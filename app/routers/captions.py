import csv
import io

from fastapi import APIRouter, Depends, File, UploadFile
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db import get_db
from app.deps import PageParams, audit, get_current_user, get_or_404, paginate
from app.errors import AppError
from app.models import ApplicationUser, Caption
from app.schemas import CaptionIn, CaptionOut, CaptionUpdate, ImportResult, Paginated

router = APIRouter(prefix="/api/captions", tags=["captions"], dependencies=[Depends(get_current_user)])
MAX_IMPORT_BYTES, MAX_ROWS = 5 * 1024 * 1024, 5000


def _key(title: str) -> str:
    return title.strip().lower()


@router.get("", response_model=Paginated[CaptionOut])
def list_captions(q: str | None = None, category: str | None = None, p: PageParams = Depends(), db: Session = Depends(get_db)):
    stmt = select(Caption).order_by(Caption.id)
    if q:
        stmt = stmt.where(Caption.title.ilike(f"%{q}%"))
    if category:
        stmt = stmt.where(Caption.category == category)
    return paginate(db, stmt, p, CaptionOut)


@router.post("", response_model=CaptionOut, status_code=201)
def create_caption(body: CaptionIn, db: Session = Depends(get_db), user: ApplicationUser = Depends(get_current_user)):
    if db.scalar(select(Caption.id).where(Caption.title_key == _key(body.title))):
        raise AppError(409, "DUPLICATE_CAPTION", "A caption with this title already exists.")
    cap = Caption(title=body.title.strip(), title_key=_key(body.title), description=body.description, category=body.category)
    db.add(cap)
    db.flush()
    audit(db, user, "create", "caption", cap.id)
    db.commit()
    return cap


@router.patch("/{caption_id}", response_model=CaptionOut)
def update_caption(caption_id: int, body: CaptionUpdate, db: Session = Depends(get_db), user: ApplicationUser = Depends(get_current_user)):
    cap = get_or_404(db, Caption, caption_id, "Caption")
    data = body.model_dump(exclude_unset=True)
    if "title" in data:
        k = _key(data["title"])
        other = db.scalar(select(Caption.id).where(Caption.title_key == k, Caption.id != caption_id))
        if other:
            raise AppError(409, "DUPLICATE_CAPTION", "A caption with this title already exists.")
        cap.title, cap.title_key = data.pop("title").strip(), k
    for k2, v in data.items():
        setattr(cap, k2, v)
    audit(db, user, "update", "caption", cap.id)
    db.commit()
    return cap


@router.delete("/{caption_id}", status_code=204)
def delete_caption(caption_id: int, db: Session = Depends(get_db), user: ApplicationUser = Depends(get_current_user)):
    cap = get_or_404(db, Caption, caption_id, "Caption")
    db.delete(cap)
    audit(db, user, "delete", "caption", caption_id)
    db.commit()


def _read_rows(name: str, raw: bytes) -> list[list[str]]:
    lower = name.lower()
    if lower.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text)) if any(c.strip() for c in row)]
    if lower.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        rows = [["" if c is None else str(c).strip() for c in r] for r in wb.active.iter_rows(values_only=True)]
        return [r for r in rows if any(r)]
    raise AppError(422, "UNSUPPORTED_FILE", "Upload a .csv or .xlsx file.")


@router.post("/import", response_model=ImportResult)
async def import_captions(dry_run: bool = True, file: UploadFile = File(...), db: Session = Depends(get_db),
                          user: ApplicationUser = Depends(get_current_user)):
    """Columns: title, description, category (optional). Use dry_run=true first to see validation results."""
    raw = await file.read(MAX_IMPORT_BYTES + 1)
    if len(raw) > MAX_IMPORT_BYTES:
        raise AppError(413, "FILE_TOO_LARGE", "Import file is larger than 5 MB.")
    try:
        rows = _read_rows(file.filename or "", raw)
    except AppError:
        raise
    except Exception:
        raise AppError(422, "UNREADABLE_FILE", "The file could not be read.")
    if rows and rows[0] and rows[0][0].lower() == "title":
        rows = rows[1:]
    if len(rows) > MAX_ROWS:
        raise AppError(413, "TOO_MANY_ROWS", f"Maximum {MAX_ROWS} rows per import.")

    seen = set(db.scalars(select(Caption.title_key)).all())
    good: list[Caption] = []
    invalid: list[str] = []
    dups = 0
    for i, r in enumerate(rows, start=1):
        title = r[0] if len(r) > 0 else ""
        desc = r[1] if len(r) > 1 else ""
        cat = (r[2] if len(r) > 2 else "") or "General"
        if not title or not desc:
            invalid.append(f"Row {i}: title and description are required.")
        elif len(title) > 200 or len(desc) > 5000 or len(cat) > 50:
            invalid.append(f"Row {i}: a value is too long.")
        elif _key(title) in seen:
            dups += 1
        else:
            seen.add(_key(title))
            good.append(Caption(title=title, title_key=_key(title), description=desc, category=cat))
    if not dry_run and good:
        db.add_all(good)
        audit(db, user, "import", "caption", len(good))
        db.commit()
    return ImportResult(dry_run=dry_run, valid=len(good), imported=0 if dry_run else len(good), duplicates=dups, invalid=invalid[:50])
